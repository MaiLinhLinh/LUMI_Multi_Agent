from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).parent
REPORT = ROOT / "Lumi_Technical_Report_Framework.docx"
DIAGRAMS = [
    ROOT / "_diagram_framework_composition.png",
    ROOT / "_diagram_plan_agent_loop.png",
    ROOT / "_diagram_runtime_loop.png",
]

NAVY = "#0B2545"
BLUE = "#2E74B5"
DARK_BLUE = "#1F4D78"
LIGHT_BLUE = "#E8EEF5"
LIGHT_GRAY = "#F2F4F7"
GREEN = "#EAF6EF"
AMBER = "#FFF4D6"
PURPLE = "#EEE7F7"
PINK = "#F9E6E8"
LINE = "#6687A8"
MUTED = "#536779"


def f(size: int, bold: bool = False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path("C:/Windows/Fonts") / name), size)


def canvas(title: str, subtitle: str):
    image = Image.new("RGB", (1800, 1320), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 110), fill=NAVY)
    draw.text((60, 23), title, font=f(35, True), fill="white")
    draw.text((60, 72), subtitle, font=f(17), fill="#D9E7F4")
    return image, draw


def box(draw, rect, title: str, lines: list[str], fill: str, *, label: str | None = None):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=18, fill=fill, outline=LINE, width=3)
    if label:
        draw.rounded_rectangle((x1 + 14, y1 - 18, x1 + 55, y1 + 20), radius=5, fill="white", outline=DARK_BLUE, width=2)
        draw.text((x1 + 34, y1 + 1), label, font=f(18, True), fill=BLUE, anchor="mm")
    draw.text(((x1 + x2) / 2, y1 + 25), title, font=f(26, True), fill=NAVY, anchor="ma")
    y = y1 + 73
    for index, line in enumerate(lines):
        is_section = line.endswith(":")
        draw.text((x1 + 28, y), line, font=f(17, is_section), fill=DARK_BLUE if is_section else NAVY)
        y += 25 if is_section else 30


def arrow(draw, start, end, text: str = "", *, color=BLUE, width=4, label_shift=(0, -22)):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        head = [(x2, y2), (x2 - 18 * direction, y2 - 10), (x2 - 18 * direction, y2 + 10)]
    else:
        direction = 1 if y2 > y1 else -1
        head = [(x2, y2), (x2 - 10, y2 - 18 * direction), (x2 + 10, y2 - 18 * direction)]
    draw.polygon(head, fill=color)
    if text:
        mx = (x1 + x2) / 2 + label_shift[0]
        my = (y1 + y2) / 2 + label_shift[1]
        tw = max(110, len(text) * 10)
        draw.rounded_rectangle((mx - tw / 2, my - 13, mx + tw / 2, my + 13), radius=5, fill="white")
        draw.text((mx, my), text, font=f(15, True), fill=DARK_BLUE, anchor="mm")


def composition():
    image, draw = canvas(
        "Framework Lumi: thành phần, nhiệm vụ và giao thức",
        "Mỗi mũi tên là một contract. Model không trực tiếp tạo DOM, URL thô, component ID hay anchor ID.",
    )
    box(draw, (70, 140, 455, 475), "Input vào Gemini Live", [
        "INPUT:", "voice/text của người dùng", "history", "Stage Map + revision hiện tại", "NHIỆM VỤ:", "đối thoại và quyết định có cần UI", "OUTPUT:", "audio/lời thoại hoặc Live tool call",
    ], AMBER, label="I")
    box(draw, (70, 565, 455, 800), "System Instructions", [
        "Gemini: cách nói, visual/state rules", "Plan: activity-first, tool loop", "Widget/Domain: contract, policy", "Không phải dữ liệu UI trực tiếp.",
    ], LIGHT_BLUE, label="S")
    box(draw, (600, 145, 1110, 535), "Gemini Live", [
        "INPUT: voice/text + Stage Map", "NHIỆM VỤ:", "• no_ui: nói trực tiếp", "• cần UI: route_request", "• panel có sẵn: present_visual", "  update_surface_state / delete_surface", "OUTPUT: PCM, lời thoại, native call",
    ], PINK, label="G")
    box(draw, (1240, 140, 1740, 535), "Live tool contract", [
        "route_request(domain_id, intent)", "present_visual(anchor_id, effect_id)", "update_surface_state(surface_id,", "  base_revision, updates)", "delete_surface(surface_id, base_revision)", "Tool response luôn do Runtime xác minh.",
    ], GREEN, label="T")
    box(draw, (600, 620, 1110, 1060), "Plan Agent", [
        "INPUT: intent, history, active surface,", "asset/template/widget/capability index", "NHIỆM VỤ:", "hiểu nhu cầu → chọn activity → đủ data", "→ kiểm tra feasibility → chọn UI", "OUTPUT (một JSON final):", "use_existing_surface_template |", "create_surface_plan | patch_surface_plan",
    ], LIGHT_BLUE, label="P")
    box(draw, (1240, 640, 1740, 1050), "Plan tool + capability server", [
        "Native tools:", "describe_widgets(widget_ids)", "describe_template(template_ids)", "call_capability(capability_id, arguments)", "Capability hiện có:", "search_web(query)", "search_image(query)", "OUTPUT: verified_data / contract detail",
    ], GREEN, label="C")
    box(draw, (70, 900, 455, 1250), "Post-processors + Browser", [
        "Compiler: widget/grid/asset/state validation", "Materializer: asset/result → source URL", "Runtime: revision + interaction verify", "Browser: render SurfaceDocument, PCM", "  và gửi select / flip / diagnostic", "OUTPUT: UI thật + feedback có cấu trúc",
    ], PURPLE, label="R")
    box(draw, (600, 1065, 1110, 1235), "SurfaceDocument", [
        "component tree + layout + props + state", "anchors + allowed effects + revision", "cùng nguồn cho browser và Stage Map",
    ], AMBER, label="D")
    # Endpoint contracts are written inside each component; arrows are deliberately
    # unlabeled so their labels never obscure the responsibilities they connect.
    arrow(draw, (455, 292), (600, 292))
    arrow(draw, (455, 682), (600, 355))
    arrow(draw, (1110, 315), (1240, 315))
    arrow(draw, (1240, 440), (1110, 440))
    arrow(draw, (855, 535), (855, 620))
    arrow(draw, (1110, 830), (1240, 830))
    arrow(draw, (1240, 930), (1110, 930))
    arrow(draw, (600, 955), (455, 1040))
    arrow(draw, (455, 1140), (600, 1160))
    arrow(draw, (855, 1065), (855, 535))
    draw.text((70, 1268), "Điểm thay thế/mở rộng: domain manifest, widget module, capability adapter, search provider, template catalog và model provider.", font=f(17, True), fill=DARK_BLUE)
    image.save(DIAGRAMS[0])


def plan_loop():
    image, draw = canvas(
        "Plan Agent: vòng lập kế hoạch trước khi sinh Surface Plan",
        "Plan Agent chỉ sinh output cuối sau khi data, widget, interaction và khả năng runtime đã đủ.",
    )
    box(draw, (70, 170, 470, 445), "Input cố định", [
        "intent từ route_request", "domain_id + history", "active surface / revision (nếu có)", "asset, template, widget index", "capability index", "verified_data + feedback trong loop",
    ], AMBER, label="1")
    box(draw, (650, 150, 1150, 590), "Reasoning nội bộ bắt buộc", [
        "1. Hiểu mục tiêu và ngữ cảnh", "2. Chọn activity/app phù hợp", "3. Liệt kê data / ảnh / state / interaction", "4. Kiểm kê catalog + widget + template", "5. Bù thiếu bằng tool cụ thể", "6. Kiểm tra mechanic có chạy thật", "7. Chỉ sau đó use / create / patch",
    ], PINK, label="2")
    box(draw, (1330, 170, 1730, 570), "Tool loop", [
        "describe_widgets", "→ props/state/action/children", "describe_template", "→ mechanic + slot + binding", "call_capability", "→ search_web / search_image", "Tool response bổ sung verified_data.",
    ], GREEN, label="3")
    box(draw, (70, 670, 470, 940), "Feasibility gate", [
        "Có facts/ảnh cần thiết?", "Widget có action/state thật?", "Frontend/backend có event thật?", "Template khớp mechanics toàn bộ?", "Nếu không: đổi mechanic hoặc route yêu cầu", "Không dựng interaction giả.",
    ], PURPLE, label="4")
    box(draw, (650, 720, 1150, 1030), "Final Surface Plan (JSON only)", [
        "A. use_existing_surface_template", "   template_id + bindings", "B. create_surface_plan", "   surface.blocks + initial_state", "C. patch_surface_plan", "   surface_id + base_revision + operations", "Không prose, Markdown hay raw URL.",
    ], LIGHT_BLUE, label="5")
    box(draw, (1330, 720, 1730, 1030), "Compiler feedback loop", [
        "invalid_widget_contract", "grid_out_of_bounds / overlap", "invalid_remote_image_result", "runtime diagnostic cần repair", "→ trả lỗi cấu trúc về Plan Agent", "→ Plan Agent sửa kế hoạch",
    ], AMBER, label="6")
    box(draw, (650, 1110, 1150, 1235), "Output hợp lệ", [
        "SurfaceDocument sau Compiler; browser payload; Stage Map/revision gửi Gemini Live.",
    ], GREEN, label="7")
    arrow(draw, (470, 305), (650, 305), "context")
    arrow(draw, (1150, 350), (1330, 350), "tool call")
    arrow(draw, (1330, 475), (1150, 475), "response")
    arrow(draw, (900, 590), (270, 670), "requirements")
    arrow(draw, (470, 805), (650, 870), "feasible")
    arrow(draw, (1150, 875), (1330, 875), "compile")
    arrow(draw, (1330, 975), (1150, 975), "structured feedback", label_shift=(0, 18))
    arrow(draw, (900, 1030), (900, 1110), "valid plan")
    image.save(DIAGRAMS[1])


def runtime_loop():
    image, draw = canvas(
        "Runtime: từ Surface Plan đến UI, Stage Map và interaction",
        "SurfaceDocument là nguồn trung gian duy nhất: browser render nó; Stage Map mô tả chính nó; Runtime giữ revision của nó.",
    )
    box(draw, (70, 180, 455, 450), "Surface Plan command", [
        "INPUT:", "create / patch / use template", "surface_id + base_revision khi patch", "NHIỆM VỤ:", "mô tả structure/state mong muốn", "OUTPUT:", "plan chưa tin cậy",
    ], LIGHT_BLUE, label="A")
    box(draw, (610, 170, 1080, 505), "Compiler", [
        "INPUT: plan + domain resources + session result", "VALIDATE: grid, widget props, children,", "state, asset, result ID, raw URL", "MATERIALIZE: catalog/search image → source", "OUTPUT: SurfaceDocument hoặc feedback",
    ], PURPLE, label="B")
    box(draw, (1250, 170, 1730, 505), "Surface Runtime", [
        "Giữ active document + revision", "Cấp component/anchor qua Compiler", "Gửi browser payload", "Gửi Stage Map context Gemini", "Verify tool calls và browser events", "State update atomic / rollback lỗi",
    ], GREEN, label="C")
    box(draw, (70, 710, 455, 1030), "Browser renderer", [
        "INPUT: SurfaceDocument client payload", "Render: web/widgets/<widget>.js", "Phát PCM; show visual effect", "OUTPUT:", "select / flip / drag-drop event", "image-load/render diagnostic",
    ], AMBER, label="D")
    box(draw, (610, 700, 1080, 1035), "SurfaceDocument + Stage Map", [
        "SurfaceDocument:", "components, layout, props, state", "anchors, effects, revision", "Stage Map:", "ASCII đúng UI rendered + anchor", "không lấy mô tả tự do từ Agent",
    ], LIGHT_BLUE, label="E")
    box(draw, (1250, 710, 1730, 1040), "Gemini Live presentation", [
        "INPUT: Stage Map + interaction verified", "TOOLS:", "present_visual(anchor, effect)", "update_surface_state(...) / delete_surface", "OUTPUT: speech/PCM + tool call", "Không tự tạo state/anchor/answer.",
    ], PINK, label="F")
    box(draw, (610, 1140, 1080, 1245), "Repair contract", [
        "Browser error → Runtime diagnostic → Plan Agent (nếu cần structural repair) → Compiler → revision mới → Stage Map mới.",
    ], GREEN, label="G")
    arrow(draw, (455, 315), (610, 315), "plan")
    arrow(draw, (1080, 340), (1250, 340), "SurfaceDocument")
    arrow(draw, (1490, 505), (300, 710), "payload to browser")
    arrow(draw, (1250, 875), (1080, 875), "Stage Map / revision")
    arrow(draw, (610, 910), (455, 910), "render snapshot")
    arrow(draw, (455, 985), (1250, 985), "interaction / diagnostic")
    arrow(draw, (1250, 785), (1080, 785), "state tool result", label_shift=(0, -20))
    arrow(draw, (845, 1035), (845, 1140), "repair needed")
    image.save(DIAGRAMS[2])


def add_run(paragraph, text: str, size: int, color: str, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def remove_old_appendix(doc):
    body = doc._element.body
    # Remove the old visual appendix and every earlier regeneration of it.
    # ASCII prefix avoids depending on the shell's legacy display encoding.
    appendix = next(
        (
            p
            for p in doc.paragraphs
            if p.style.name == "Heading 1" and p.text.strip().startswith("13.")
        ),
        None,
    )
    if appendix is None:
        return
    remove = False
    for child in list(body):
        if child is appendix._p:
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def update_doc():
    composition()
    plan_loop()
    runtime_loop()
    doc = Document(REPORT)
    remove_old_appendix(doc)
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.style = "Heading 1"
    add_run(heading, "13. S\u01a1 \u0111\u1ed3 framework chi ti\u1ebft", 16, "2E74B5", True)
    p = doc.add_paragraph()
    add_run(p, "Ba sơ đồ dưới đây dùng cùng cách phân ranh giới với hệ thống tham chiếu: input, system instructions, LLM/tool loop, post-processors, output browser và feedback. Điểm khác là Lumi sinh Surface Plan/SurfaceDocument thay vì HTML/CSS/JS tự do.", 10.5, "0B2545")
    captions = [
        "Hình 1. Thành phần framework, nhiệm vụ từng thành phần và các native tool/capability hiện có.",
        "Hình 2. Vòng Plan Agent: input, reasoning, tool loop, feasibility gate, output và compiler repair.",
        "Hình 3. Runtime contract: Compiler, SurfaceDocument, browser, Stage Map, Gemini Live và state/repair loop.",
    ]
    for index, (path, caption_text) in enumerate(zip(DIAGRAMS, captions, strict=True)):
        if index:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(6.45))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(cap, caption_text, 9, "536779", True)
    doc.save(REPORT)


if __name__ == "__main__":
    update_doc()
